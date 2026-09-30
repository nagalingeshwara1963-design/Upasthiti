"""Portal/review authorization tests use temporary SQLite databases only."""
from datetime import date, datetime, time, timedelta
import tempfile
import time as clock
import unittest
from pathlib import Path
from unittest.mock import patch

from app import auth, chat_assistant, recovery, stats
from app.db import DB
from app.portal_service import StudentPortalService
from app.review_requests import ReviewRequestService
from app import reportstore
from app import reports, exports
from app import demo_mode
from app.email_service import EmailService, EmailConfigurationError
from app import backup
from openpyxl import load_workbook


class PortalFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="upasthiti-portal-")
        self.ensure_patch = patch("app.db.config.ensure_dirs"); self.ensure_patch.start()
        self.db = DB(Path(self.temp.name) / "test.db")
        self.class_a = self.db.add_class("Class A")
        self.class_b = self.db.add_class("Class B")
        self.a = self.db.save_student(self.class_a, "A-1", "A Student", [])
        self.b = self.db.save_student(self.class_a, "B-1", "B Student", [])
        self.other = self.db.save_student(self.class_b, "C-1", "C Student", [])

    def tearDown(self):
        auth.logout(); self.db.con.close(); self.ensure_patch.stop(); self.temp.cleanup()

    def session(self, class_id, student_id, status="A", day=None):
        return self.db.save_session(class_id, {"date": (day or date.today()).isoformat(), "time": "09:00",
            "faculty": "Faculty", "subject": "Computing", "period": "1", "mode": "fast"},
            [{"student_id": student_id, "status": status}])

    def student_login(self, sid, name, code):
        self.assertEqual(auth.student_login(self.db, name, code, code), "student")
        self.assertEqual(auth.current_student_id(), sid)
        auth.change_student_password(self.db, "a-strong-new-password")


class RecoveryTests(PortalFixture):
    def test_archive_preserves_attendance_edits_reviews_and_historical_reports(self):
        sid = self.session(self.class_a, self.a, "P")
        attendance_id = self.db.q("SELECT id FROM attendance WHERE session_id=? AND student_id=?", (sid, self.a))[0]["id"]
        self.db.change_status(sid, self.a, "A", by_user="admin")
        self.db.create_review_request(self.a, sid, attendance_id, "P", "mismatch", "synthetic test")
        self.db.archive_student(self.class_a, "A-1")

        self.assertEqual([s["code"] for s in self.db.students(self.class_a)], ["B-1"])
        archived = self.db.students(self.class_a, include_archived=True)
        self.assertEqual(next(s for s in archived if s["id"] == self.a)["active"], 0)
        self.assertEqual(self.db.session_rows(sid)[0]["status"], "A")
        self.assertEqual(len(self.db.edits_for(sid)), 1)
        self.assertEqual(len(self.db.list_review_requests()), 1)
        rows, _held = stats.range_students(self.db, self.class_a,
            date.today().isoformat(), date.today().isoformat(), include_archived=True)
        self.assertEqual(next(r for r in rows if r["id"] == self.a)["absent"], 1)
        active_rows, _held = stats.range_students(self.db, self.class_a,
            date.today().isoformat(), date.today().isoformat())
        self.assertNotIn(self.a, [r["id"] for r in active_rows])
        workbook_path = Path(self.temp.name) / "archived-history.xlsx"
        reports.range_xlsx(self.db, self.class_a, date.today().isoformat(),
                           date.today().isoformat(), workbook_path)
        workbook_rows = list(load_workbook(workbook_path, data_only=True).active.iter_rows(values_only=True))
        self.assertIn(("A-1", "A Student", 0, 1, 0, 1, 0.0), workbook_rows)
        self.assertIsNone(auth.student_login(self.db, "A Student", "A-1", "A-1"))

    def test_old_student_schema_migrates_without_losing_existing_rows(self):
        import sqlite3
        path = Path(self.temp.name) / "old-schema.db"
        con = sqlite3.connect(path)
        con.executescript("""
            CREATE TABLE classes(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL,
                college TEXT DEFAULT '', year TEXT DEFAULT '', branch TEXT DEFAULT '', section TEXT DEFAULT '');
            CREATE TABLE students(id INTEGER PRIMARY KEY, class_id INTEGER NOT NULL, code TEXT NOT NULL,
                name TEXT DEFAULT '', UNIQUE(class_id, code));
            INSERT INTO classes(id,name) VALUES(7,'Legacy class');
            INSERT INTO students(id,class_id,code,name) VALUES(9,7,'LEG-1','Legacy student');
        """)
        con.commit(); con.close()
        migrated = DB(path)
        try:
            row = migrated.q("SELECT * FROM students WHERE id=9")[0]
            self.assertEqual((row["class_id"], row["code"], row["name"], row["active"]),
                             (7, "LEG-1", "Legacy student", 1))
        finally:
            migrated.con.close()

    def test_math_empty_boundary_and_impossible_target(self):
        empty = recovery.plan(0, 0, 75)
        self.assertIsNone(empty["current"]); self.assertIsNone(empty["classes_needed"])
        self.assertIsNone(empty["safe_to_miss"])
        self.assertEqual(recovery.plan(2, 2, 75)["classes_needed"], 4)
        self.assertIsNone(recovery.plan(3, 1, 100)["classes_needed"])
        self.assertEqual(stats.counts(["P", "L", "A", "E", "OD"])["pct"], 66.66666666666667)

    def test_student_identity_scopes_summary_report_and_history(self):
        self.session(self.class_a, self.a, "P")
        self.session(self.class_a, self.b, "A")
        self.student_login(self.a, "A Student", "A-1")
        portal = StudentPortalService(self.db)
        self.assertEqual(portal.summary()["student"]["id"], self.a)
        self.assertEqual(len(portal.history()), 1)
        report = portal.report_summary(date.today().isoformat(), date.today().isoformat())
        self.assertEqual(report["attended"], 1); self.assertEqual(len(report["records"]), 1)
        self.assertNotIn("B Student", repr(portal.notifications()))
        self.assertNotIn("B Student", repr(chat_assistant.LocalAssistant(self.db).answer("attendance")))

    def test_initial_password_forces_change_old_credential_rejected(self):
        self.assertEqual(auth.student_login(self.db, "A Student", "A-1", "A-1"), "student")
        self.assertTrue(auth.student_must_change_password())
        with self.assertRaises(PermissionError): StudentPortalService(self.db).summary()
        auth.change_student_password(self.db, "a-strong-new-password")
        auth.logout()
        self.assertIsNone(auth.student_login(self.db, "A Student", "A-1", "A-1"))
        self.assertEqual(auth.student_login(self.db, "A Student", "A-1", "a-strong-new-password"), "student")

    def test_unicode_student_name_login_is_exact_and_casefolded(self):
        student = self.db.save_student(self.class_a, "U-1", "Élodie O’Neil", [])
        self.assertEqual(auth.student_login(self.db, "élodie o’neil", "U-1", "U-1"), "student")
        self.assertEqual(auth.current_student_id(), student)

    def test_review_student_cannot_request_other_students_session(self):
        other_session = self.session(self.class_a, self.b, "A")
        self.student_login(self.a, "A Student", "A-1")
        with self.assertRaises(PermissionError): ReviewRequestService(self.db).submit(other_session, "I was present")

    def test_review_window_exact_10_day_success_and_failed_notification(self):
        old_day = date.today() - timedelta(days=31)
        sid = self.session(self.class_a, self.a, "A", old_day)
        delivery = self.db.create_email_delivery({"kind": "absence", "status": "sent", "dedupe_key": "sent",
            "student_id": self.a, "session_id": sid, "class_id": self.class_a})
        now = clock.time()
        self.db.con.execute("UPDATE email_deliveries SET updated_at=? WHERE id=?", (now - 10 * 86400, delivery["id"])); self.db.con.commit()
        self.student_login(self.a, "A Student", "A-1")
        req, made = ReviewRequestService(self.db).submit(sid, "I was present", now=now)
        self.assertTrue(made); self.assertEqual(req["status"], "Escalated")
        sid2 = self.session(self.class_a, self.a, "A", old_day)
        self.db.create_email_delivery({"kind": "absence", "status": "failed", "dedupe_key": "failed",
            "student_id": self.a, "session_id": sid2, "class_id": self.class_a})
        with self.assertRaises(PermissionError): ReviewRequestService(self.db).submit(sid2, "Other", now=now)

    def test_review_report_must_cover_session_and_be_successful(self):
        old_day = date.today() - timedelta(days=40)
        sid = self.session(self.class_a, self.a, "A", old_day)
        self.db.create_email_delivery({"kind": "student_report", "status": "sent", "dedupe_key": "report-outside",
            "student_id": self.a, "class_id": self.class_a, "report_from": date.today().isoformat(), "report_to": date.today().isoformat()})
        self.student_login(self.a, "A Student", "A-1")
        with self.assertRaises(PermissionError): ReviewRequestService(self.db).submit(sid, "Other")
        self.db.con.execute("UPDATE email_deliveries SET report_from=?,report_to=?,updated_at=? WHERE dedupe_key='report-outside'",
            (old_day.isoformat(), old_day.isoformat(), clock.time() - 86400)); self.db.con.commit()
        req, made = ReviewRequestService(self.db).submit(sid, "Other")
        self.assertTrue(made); self.assertEqual(req["status"], "Escalated")

    def test_review_lifecycle_requires_human_action_and_audits_change(self):
        session = self.session(self.class_a, self.a, "A")
        self.student_login(self.a, "A Student", "A-1")
        request, created = ReviewRequestService(self.db).submit(session, "I was present", "I attended the class.")
        self.assertTrue(created)
        duplicate, created_again = ReviewRequestService(self.db).submit(session, "I was present")
        self.assertFalse(created_again); self.assertEqual(duplicate["id"], request["id"])
        auth.logout()
        auth.set_password(self.db, "faculty", "reviewer", "reviewer-test-password")
        self.db.set_faculty_classes("reviewer", [self.class_a])
        auth.login(self.db, "reviewer", "reviewer-test-password")
        service = ReviewRequestService(self.db)
        self.assertEqual(service.transition(request["id"], "start")["status"], "Under Review")
        resolved = service.transition(request["id"], "change", "Evidence confirmed presence.", "P")
        self.assertEqual(resolved["status"], "Resolved — Changed")
        edit = self.db.edits_for(session)[0]
        self.assertEqual((edit["old"], edit["new"], edit["by_user"]), ("A", "P", "faculty:reviewer"))
        self.assertEqual(len(self.db.review_events(request["id"])), 3)
        auth.logout()
        self.assertEqual(auth.student_login(self.db, "A Student", "A-1", "a-strong-new-password"), "student")
        self.assertEqual(StudentPortalService(self.db).review_requests()[0]["status"], "Resolved — Changed")

    def test_faculty_class_and_edit_window_boundaries(self):
        auth.set_password(self.db, "faculty", "fac", "test-password")
        self.db.set_faculty_classes("fac", [self.class_a])
        self.assertEqual(auth.login(self.db, "fac", "test-password"), "faculty")
        self.assertTrue(auth.can_access_class(self.db, self.class_a))
        self.assertFalse(auth.can_access_class(self.db, self.class_b))
        today = date.today()
        for age, expected in ((29, True), (30, True), (31, False)):
            with self.subTest(age=age):
                self.assertEqual(auth.can_edit_session(session_date=(today - timedelta(days=age)).isoformat(),
                    db=self.db, class_id=self.class_a), expected)
        session_b = self.session(self.class_b, self.other, "A")
        attendance_id = self.db.q("SELECT id FROM attendance WHERE session_id=?", (session_b,))[0]["id"]
        req, _ = self.db.create_review_request(self.other, session_b, attendance_id, "A", "Other", "")
        with self.assertRaises(PermissionError): ReviewRequestService(self.db).transition(req["id"], "keep")
        auth.set_password(self.db, "admin", "admin", "test-admin-password")
        self.assertEqual(auth.login(self.db, "admin", "wrong"), None)
        self.assertIsNone(auth.current_role())

    def test_retention_removes_only_old_evidence_and_keeps_active_review(self):
        with tempfile.TemporaryDirectory(prefix="upasthiti-evidence-") as root, \
             patch("app.reportstore.config.REPORTS_DIR", Path(root)):
            old_day = date.today() - timedelta(days=31)
            old_session = self.session(self.class_a, self.a, "A", old_day)
            active_session = self.session(self.class_a, self.b, "A", old_day)
            for sid in (old_session, active_session):
                folder = Path(root) / f"session_{sid}"; folder.mkdir()
                (folder / "photo_1.jpg").write_bytes(b"evidence")
                (folder / "result.json").write_text("{}", encoding="utf-8")
            attendance_id = self.db.q("SELECT id FROM attendance WHERE session_id=?", (active_session,))[0]["id"]
            self.db.create_review_request(self.b, active_session, attendance_id, "A", "Other", "")
            self.db.put("photo_retention_days", 30)
            removed = reportstore.purge_expired_session_artifacts(self.db)
            self.assertEqual(removed, 2)
            self.assertFalse((Path(root) / f"session_{old_session}" / "photo_1.jpg").exists())
            self.assertTrue((Path(root) / f"session_{active_session}" / "photo_1.jpg").exists())
            self.assertEqual(self.db.q("SELECT COUNT(*) n FROM attendance")[0]["n"], 2)

    def test_spreadsheet_exports_treat_user_text_as_text(self):
        malicious_class = self.db.add_class("=1+1")
        student = self.db.save_student(malicious_class, "+CMD|' /C calc'!A0", "=HYPERLINK(\"https://example.invalid\")", [])
        session = self.session(malicious_class, student, "A")
        with tempfile.TemporaryDirectory(prefix="upasthiti-export-") as root:
            range_path = Path(root) / "range.xlsx"
            reports.range_xlsx(self.db, malicious_class, date.today().isoformat(), date.today().isoformat(), range_path)
            sheet = load_workbook(range_path, data_only=False)["Attendance"]
            self.assertEqual(sheet["B2"].value, "=1+1"); self.assertEqual(sheet["B2"].data_type, "s")
            self.assertEqual(sheet["A5"].data_type, "s"); self.assertEqual(sheet["B5"].data_type, "s")
            self.db.change_status(session, student, "P", "=staff")
            audit_path = Path(root) / "audit.xlsx"; exports.audit_to_xlsx(self.db, audit_path)
            audit = load_workbook(audit_path, data_only=False)["Audit log"]
            self.assertEqual(audit["B2"].data_type, "s"); self.assertEqual(audit["J2"].data_type, "s")

    def test_demo_database_is_synthetic_and_email_is_hard_disabled(self):
        with tempfile.TemporaryDirectory(prefix="upasthiti-demo-test-") as root:
            old_paths = {key: getattr(__import__("app.config", fromlist=["config"]), key)
                         for key in demo_mode._PATH_KEYS}
            db = None; mail = None
            try:
                db, ids = demo_mode.create_demo_database(root)
                self.assertEqual(Path(__import__("app.config", fromlist=["config"]).DB_PATH).parent,
                                 Path(root).resolve() / "data")
                self.assertEqual(db.q("SELECT COUNT(*) n FROM sessions")[0]["n"], 3)
                self.assertTrue(all(s["code"].startswith("DEMO") for s in db.q("SELECT code FROM students")))
                self.assertEqual(auth.login(db, "admin", "demo-admin-only"), "admin")
                mail = EmailService(db, provider=object(), disabled=True)
                with self.assertRaises(EmailConfigurationError): mail.send_test_email("person@example.edu")
                self.assertEqual(db.q("SELECT COUNT(*) n FROM email_deliveries")[0]["n"], 0)
            finally:
                if mail:
                    mail.shutdown(timeout=1); mail._thread.join(timeout=2)
                if db: db.con.close()
                demo_mode.restore_config_paths(old_paths)

    def test_backup_validation_and_restore_are_atomic_and_complete(self):
        with tempfile.TemporaryDirectory(prefix="upasthiti-backup-test-") as root:
            data = Path(root) / "data"; data.mkdir()
            patches = [patch("app.backup.config.DATA", data), patch("app.backup.config.DB_PATH", data / "upasthiti.db"),
                       patch("app.backup.config.BACKUP_DIR", data / "backups"), patch("app.backup.config.ensure_dirs", lambda: [p.mkdir(parents=True, exist_ok=True) for p in (data, data / "backups")])]
            for p in patches: p.start()
            try:
                for folder in backup.ASSETS:
                    (data / folder).mkdir(parents=True, exist_ok=True)
                (data / "enroll" / "sample.jpg").write_bytes(b"sample enrollment image")
                source = Path(root) / "backup.zip"
                backup.create_backup(self.db, source)
                self.assertTrue(backup.validate_backup(source)["valid"])
                self.db.add_class("Added after backup")
                (data / "enroll" / "sample.jpg").write_bytes(b"changed")
                restored = backup.restore_backup(self.db, source)
                self.db = restored
                self.assertFalse(self.db.q("SELECT 1 FROM classes WHERE name='Added after backup'"))
                self.assertEqual((data / "enroll" / "sample.jpg").read_bytes(), b"sample enrollment image")
                self.assertTrue(list((data / "backups").glob("pre_restore_*.zip")))
            finally:
                try: self.db.con.close()
                except Exception: pass
                for p in reversed(patches): p.stop()


if __name__ == "__main__": unittest.main()
