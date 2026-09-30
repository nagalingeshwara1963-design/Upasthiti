"""Small standalone exports: the audit log (every attendance edit, across all classes)."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
from . import stats

HEADERS = ["When changed", "Class", "Session date", "Subject", "Period", "Student ID", "Name",
          "Old status", "New status", "Changed by"]


def _spreadsheet_text(value):
    text = "" if value is None else str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _rows(db):
    out = []
    for e in db.all_edits():
        when = datetime.fromtimestamp(e["at"]).strftime("%Y-%m-%d %H:%M")
        out.append([when, e["class_name"], e["date"], e["subject"] or "", e["period"] or "", e["code"], e["name"] or "",
                    stats.LABEL.get(e["old"], e["old"]), stats.LABEL.get(e["new"], e["new"]), e["by_user"]])
    return out


def audit_to_xlsx(db, path):
    rows = _rows(db)
    wb = Workbook(); ws = wb.active; ws.title = "Audit log"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    for r in rows:
        row_no = ws.max_row + 1
        for i, value in enumerate(r, 1):
            ws.cell(row=row_no, column=i, value=_spreadsheet_text(value))
    for i, w in enumerate([17, 20, 13, 14, 8, 14, 20, 11, 11, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if not rows:
        ws.append(["No edits have been made yet."])
    wb.save(str(path)); return str(path)


def audit_to_csv(db, path):
    rows = _rows(db)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(HEADERS)
        for r in rows: w.writerow([_spreadsheet_text(v) for v in r])
        if not rows: w.writerow(["No edits have been made yet."])
    return str(path)
