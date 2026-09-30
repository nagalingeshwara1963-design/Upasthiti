"""Word and Excel report builders: one session, one class-day, one class-month."""
import calendar, io, json
from datetime import datetime
from pathlib import Path
import cv2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from . import config, stats, reportstore

R = {  # report wording (en / kn / hi) - have a native speaker review kn and hi before release
 "title": ("Attendance Report", "ಹಾಜರಾತಿ ವರದಿ", "उपस्थिति रिपोर्ट"),
 "day_title": ("Daily Attendance Report", "ದೈನಿಕ ಹಾಜರಾತಿ ವರದಿ", "दैनिक उपस्थिति रिपोर्ट"),
 "month_title": ("Monthly Attendance Summary", "ಮಾಸಿಕ ಹಾಜರಾತಿ ಸಾರಾಂಶ", "मासिक उपस्थिति सारांश"),
 "class": ("Class", "ತರಗತಿ", "कक्षा"), "date": ("Date", "ದಿನಾಂಕ", "तारीख"), "time": ("Time", "ಸಮಯ", "समय"),
 "faculty": ("Faculty", "ಅಧ್ಯಾಪಕ", "शिक्षक"), "subject": ("Subject", "ವಿಷಯ", "विषय"), "period": ("Period", "ಅವಧಿ", "पीरियड"),
 "month": ("Month", "ತಿಂಗಳು", "महीना"),
 "present": ("Present", "ಹಾಜರು", "उपस्थित"), "absent": ("Absent", "ಗೈರು", "अनुपस्थित"), "late": ("Late", "ತಡ", "देर से"),
 "excused": ("Excused", "ರಿಯಾಯಿತಿ", "क्षमा"), "leave": ("Leave / OD", "ರಜೆ / ಕರ್ತವ್ಯ", "अवकाश / ड्यूटी"),
 "attendance": ("Attendance", "ಹಾಜರಾತಿ", "उपस्थिति"),
 "photos": ("Group photos", "ಗುಂಪು ಫೋಟೋಗಳು", "ग्रुप फोटो"),
 "legend": ("Green box = matched, amber = uncertain, grey = not matched. Labels show the last 4 characters of the student ID.",
            "ಹಸಿರು = ಹೊಂದಾಣಿಕೆ, ಕಿತ್ತಳೆ = ಅನಿಶ್ಚಿತ, ಬೂದು = ಹೊಂದಿಲ್ಲ. ಲೇಬಲ್‌ಗಳು ವಿದ್ಯಾರ್ಥಿ ಐಡಿಯ ಕೊನೆಯ 4 ಅಕ್ಷರಗಳು.",
            "हरा बॉक्स = मिलान, नारंगी = अनिश्चित, धूसर = मिलान नहीं। लेबल छात्र आईडी के अंतिम 4 अक्षर हैं।"),
 "list": ("Attendance list", "ಹಾಜರಾತಿ ಪಟ್ಟಿ", "उपस्थिति सूची"),
 "id": ("Student ID", "ವಿದ್ಯಾರ್ಥಿ ಐಡಿ", "छात्र आईडी"), "name": ("Name", "ಹೆಸರು", "नाम"), "status": ("Status", "ಸ್ಥಿತಿ", "स्थिति"),
 "conf": ("Confidence", "ವಿಶ್ವಾಸ", "विश्वास"),
 "absent_photos": ("Absent students (enrollment photos)", "ಗೈರು ವಿದ್ಯಾರ್ಥಿಗಳು (ನೋಂದಣಿ ಫೋಟೋ)", "अनुपस्थित छात्र (नामांकन फोटो)"),
 "changes": ("Manual changes and uncertain marks", "ಕೈಯಾರೆ ಬದಲಾವಣೆಗಳು ಮತ್ತು ಅನಿಶ್ಚಿತ ಗುರುತುಗಳು", "मैन्युअल बदलाव और अनिश्चित निशान"),
 "system": ("System result", "ವ್ಯವಸ್ಥೆಯ ಫಲಿತಾಂಶ", "सिस्टम परिणाम"), "final": ("Final", "ಅಂತಿಮ", "अंतिम"), "note": ("Note", "ಟಿಪ್ಪಣಿ", "टिप्पणी"),
 "student": ("Student", "ವಿದ್ಯಾರ್ಥಿ", "छात्र"),
 "none": ("None", "ಯಾವುದೂ ಇಲ್ಲ", "कोई नहीं"),
 "signature": ("Faculty signature", "ಅಧ್ಯಾಪಕರ ಸಹಿ", "शिक्षक हस्ताक्षर"),
 "sessions": ("Sessions", "ಅವಧಿಗಳು", "सत्र"), "total": ("Total", "ಒಟ್ಟು", "कुल"),
 "pct": ("Attendance %", "ಹಾಜರಾತಿ %", "उपस्थिति %"),
 "ok": ("OK", "ಸರಿ", "ठीक"), "warn": ("Warning", "ಎಚ್ಚರಿಕೆ", "चेतावनी"), "crit": ("Below limit", "ಮಿತಿಗಿಂತ ಕಡಿಮೆ", "सीमा से कम"),
 "rule": ("Present and Late count as attended. Excused and Leave/OD are not counted against the student.",
          "ಹಾಜರು ಮತ್ತು ತಡ ಹಾಜರಾತಿ ಎಂದು ಪರಿಗಣಿಸಲಾಗುತ್ತದೆ. ರಿಯಾಯಿತಿ ಮತ್ತು ರಜೆ/ಕರ್ತವ್ಯ ವಿದ್ಯಾರ್ಥಿಯ ವಿರುದ್ಧ ಲೆಕ್ಕಿಸುವುದಿಲ್ಲ.",
          "उपस्थित और देर से को उपस्थिति माना जाता है। क्षमा और अवकाश/ड्यूटी छात्र के विरुद्ध नहीं गिने जाते।"),
 "flag_weak": ("weak match", "ದುರ್ಬಲ ಹೊಂದಾಣಿಕೆ", "कमजोर मिलान"), "flag_close": ("similar to another student", "ಇನ್ನೊಬ್ಬ ವಿದ್ಯಾರ್ಥಿಯಂತೆ ಕಾಣುತ್ತದೆ", "दूसरे छात्र जैसा"),
 "flag_models": ("face models disagreed", "ಮಾದರಿಗಳು ಒಪ್ಪಲಿಲ್ಲ", "मॉडल असहमत"),
 "edited_later": ("Edited after saving", "ಉಳಿಸಿದ ನಂತರ ಬದಲಾಯಿಸಲಾಗಿದೆ", "सहेजने के बाद बदला गया"),
}
LANG_IDX = {"en": 0, "kn": 1, "hi": 2}
FLAG_KEY = {"weak_match": "flag_weak", "close_call": "flag_close", "models_disagree": "flag_models"}
STATUS_KEY = {"P": "present", "A": "absent", "L": "late", "E": "excused", "OD": "leave"}
STATUS_HEX = {"P": "2E7D32", "L": "EF6C00", "A": "C62828", "E": "5C6370", "OD": "5C6370"}
FILL = {"P": "E8F5E9", "L": "FFF3E0", "A": "FDECEA", "E": "ECEFF1", "OD": "ECEFF1"}
LEVEL_HEX = {"ok": "2E7D32", "warn": "EF6C00", "critical": "C62828"}


class Rep:
    def __init__(self, lang="en"):
        self.i = LANG_IDX.get(lang, 0); self.lang = lang
        self.font = "Nirmala UI" if lang in ("kn", "hi") else "Arial"
    def t(self, k): return R[k][self.i]
    def st(self, code): return R[STATUS_KEY.get(code, "absent")][self.i]
    def flag_text(self, flags):
        return ", ".join(self.t(FLAG_KEY[f]) for f in (flags or []) if f in FLAG_KEY)


def _doc(rep):
    d = Document()
    sec = d.sections[0]; sec.left_margin = sec.right_margin = Inches(0.8); sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = d.styles["Normal"]; st.font.name = rep.font; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:cs"), rep.font); st.element.rPr.rFonts.set(qn("w:eastAsia"), rep.font)
    return d


def _run(p, text, rep, bold=False, size=None, color=None, italic=False):
    r = p.add_run(text); r.bold = bold; r.italic = italic
    r.font.name = rep.font; r._element.rPr.rFonts.set(qn("w:cs"), rep.font)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r


def _head(d, rep, db, title):
    college = db.get("college_name", "")
    logo = db.get("logo_path", "")
    if logo and Path(logo).exists():
        try: d.add_picture(logo, height=Inches(0.6))
        except Exception: pass
    if college:
        p = d.add_paragraph(); _run(p, college, rep, bold=True, size=13)
    p = d.add_paragraph(); _run(p, title, rep, bold=True, size=18, color="0D47A1")


def _table(d, rep, headers, rows, widths=None, colour_col=None, colour_fn=None):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; _run(c.paragraphs[0], h, rep, bold=True, size=9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            col = colour_fn(row, i) if colour_fn else None
            _run(cells[i].paragraphs[0], str(v), rep, size=9.5, color=col, bold=bool(col))
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths): r.cells[i].width = Inches(w)
    return t


def _info_table(d, rep, pairs):
    t = d.add_table(rows=0, cols=4); t.style = "Table Grid"
    for i in range(0, len(pairs), 2):
        cells = t.add_row().cells
        for j in range(2):
            if i + j < len(pairs):
                k, v = pairs[i + j]
                cells[2 * j].text = ""; _run(cells[2 * j].paragraphs[0], k, rep, bold=True, size=9.5)
                cells[2 * j + 1].text = ""; _run(cells[2 * j + 1].paragraphs[0], str(v), rep, size=9.5)


def _signature(d, rep):
    d.add_paragraph()
    p = d.add_paragraph(); _run(p, f"{rep.t('signature')}: ______________________        {rep.t('date')}: ______________", rep)


def _enroll_thumb(db, student_id, size=110):
    r = db.q("SELECT path FROM enroll_photos WHERE student_id=? ORDER BY id LIMIT 1", (student_id,))
    if not r: return None
    img = cv2.imread(r[0]["path"])
    if img is None: return None
    h, w = img.shape[:2]; s = size / max(h, w)
    ok, buf = cv2.imencode(".jpg", cv2.resize(img, (int(w * s), int(h * s)), interpolation=cv2.INTER_AREA))
    return io.BytesIO(buf.tobytes()) if ok else None


def _change_rows(db, rep, sid, rows):
    """Rows for the 'manual changes and uncertain marks' section."""
    out = []; edited = {e["student_id"] for e in db.edits_for(sid)}
    for r in rows:
        auto = r.get("auto_status") or r["status"]
        flags = rep.flag_text((r.get("flags") or "").split(",")) if r.get("flags") else ""
        notes = [x for x in (flags, rep.t("edited_later") if r["student_id"] in edited else "") if x]
        if auto != r["status"] or flags or r["student_id"] in edited:
            out.append([f"{r['name'] or r['code']} ({r['code']})", rep.st(auto), rep.st(r["status"]), "; ".join(notes)])
    return out


# =========================================================================== SESSION
def session_docx(db, sid, path, lang="en"):
    rep = Rep(lang); s = db.session(sid); rows = db.session_rows(sid)
    d = _doc(rep); _head(d, rep, db, rep.t("title"))
    k = stats.counts([r["status"] for r in rows])
    _info_table(d, rep, [(rep.t("class"), s["class_name"]), (rep.t("date"), s["date"]), (rep.t("time"), s["time"]),
                         (rep.t("faculty"), s["faculty"] or "-"), (rep.t("subject"), s["subject"] or "-"), (rep.t("period"), s["period"] or "-")])
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    cnt = {c: sum(1 for r in rows if r["status"] == c) for c in STATUS_KEY}
    pct_txt = "-" if k["pct"] is None else f"{k['pct']:.1f}%"
    _run(p, f"{rep.t('present')}: {cnt['P']}   {rep.t('late')}: {cnt['L']}   {rep.t('absent')}: {cnt['A']}   {rep.t('excused')}: {cnt['E']}   {rep.t('leave')}: {cnt['OD']}   |   "
            f"{rep.t('pct')}: {pct_txt}", rep, bold=True, size=11)
    bundle = reportstore.load(sid)
    if bundle["images"]:
        h = d.add_paragraph(); _run(h, rep.t("photos"), rep, bold=True, size=13, color="0D47A1")
        for i, ip in enumerate(bundle["images"]):
            try: d.add_picture(ip, width=Inches(6.2))
            except Exception: continue
            label = (bundle["meta"].get("photos") or [{}] * (i + 1))[i].get("label", "") if bundle["meta"].get("photos") else ""
            if label:
                c = d.add_paragraph(); _run(c, label, rep, italic=True, size=9)
        c = d.add_paragraph(); _run(c, rep.t("legend"), rep, italic=True, size=8.5, color="5C6370")
    h = d.add_paragraph(); _run(h, rep.t("list"), rep, bold=True, size=13, color="0D47A1")
    body = [[n, r["code"], r["name"] or "", rep.st(r["status"]), "-" if r["status"] != "P" and not r["confidence"] else f"{r['confidence']:.0f}%"]
            for n, r in enumerate(rows, 1)]
    _table(d, rep, ["#", rep.t("id"), rep.t("name"), rep.t("status"), rep.t("conf")], body, [0.4, 1.7, 2.4, 1.2, 1.0],
           colour_fn=lambda row, i: STATUS_HEX.get(next((c for c, kk in STATUS_KEY.items() if rep.st(c) == row[3]), "P")) if i == 3 else None)
    ab = [r for r in rows if r["status"] == "A"]
    if ab:
        h = d.add_paragraph(); h.paragraph_format.space_before = Pt(10); _run(h, rep.t("absent_photos"), rep, bold=True, size=13, color="0D47A1")
        cols = 5; t = d.add_table(rows=0, cols=cols)
        for i in range(0, len(ab), cols):
            cells = t.add_row().cells
            for j, r in enumerate(ab[i:i + cols]):
                cells[j].text = ""; par = cells[j].paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                th = _enroll_thumb(db, r["student_id"])
                if th: par.add_run().add_picture(th, height=Inches(0.95))
                q = cells[j].add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; _run(q, r["name"] or r["code"], rep, bold=True, size=8.5)
                q2 = cells[j].add_paragraph(); q2.alignment = WD_ALIGN_PARAGRAPH.CENTER; _run(q2, r["code"], rep, size=7.5, color="5C6370")
    h = d.add_paragraph(); h.paragraph_format.space_before = Pt(10); _run(h, rep.t("changes"), rep, bold=True, size=13, color="0D47A1")
    ch = _change_rows(db, rep, sid, rows)
    if ch: _table(d, rep, [rep.t("student"), rep.t("system"), rep.t("final"), rep.t("note")], ch, [2.4, 1.2, 1.2, 2.0])
    else: _run(d.add_paragraph(), rep.t("none"), rep, size=10)
    _signature(d, rep)
    d.save(str(path)); return str(path)



def _xl_head(ws, headers, row=1):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row=row, column=c, value=h); x.font = Font(name="Arial", bold=True); x.fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
        x.data_type = "s"
        x.alignment = Alignment(wrap_text=True, vertical="top")


def _xl_text(cell, value):
    """Store user-controlled text as text, even if it begins with a formula marker."""
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    return cell


def session_xlsx(db, sid, path, lang="en"):
    rep = Rep(lang); s = db.session(sid); rows = db.session_rows(sid)
    wb = Workbook(); sm = wb.active; sm.title = "Summary"; at = wb.create_sheet("Attendance")
    f = lambda **k: Font(name=rep.font, **k)
    _xl_head(at, ["#", rep.t("id"), rep.t("name"), rep.t("status"), "Code", rep.t("conf") + " %", "Distance", rep.t("note")])
    for n, r in enumerate(rows, 1):
        note = rep.flag_text((r.get("flags") or "").split(",")) if r.get("flags") else ""
        if (r.get("auto_status") or r["status"]) != r["status"]: note = (note + "; " if note else "") + f"{rep.st(r['auto_status'])} → {rep.st(r['status'])}"
        vals = [n, r["code"], r["name"] or "", rep.st(r["status"]), r["status"], None if not r["confidence"] else round(r["confidence"], 1),
                None if r["dist"] is None else round(r["dist"], 4), note]
        for c, v in enumerate(vals, 1):
            x = at.cell(row=n + 1, column=c)
            if c in (2, 3): _xl_text(x, v)
            else: x.value = v
            x.font = f()
        at.cell(row=n + 1, column=4).fill = PatternFill("solid", start_color=FILL.get(r["status"], "FFFFFF"), end_color=FILL.get(r["status"], "FFFFFF"))
    last = len(rows) + 1
    for i, w in enumerate([5, 18, 28, 14, 7, 12, 10, 44], 1): at.column_dimensions[get_column_letter(i)].width = w
    at.freeze_panes = "A2"
    sm["A1"] = rep.t("title"); sm["A1"].font = f(bold=True, size=14)
    info = [(rep.t("class"), s["class_name"]), (rep.t("date"), s["date"]), (rep.t("time"), s["time"]), (rep.t("faculty"), s["faculty"]),
            (rep.t("subject"), s["subject"]), (rep.t("period"), s["period"])]
    for i, (k, v) in enumerate(info, 3):
        sm.cell(row=i, column=1, value=k).font = f(bold=True); _xl_text(sm.cell(row=i, column=2), v).font = f()
    r0 = 10
    for j, (code, key) in enumerate((("P", "present"), ("L", "late"), ("A", "absent"), ("E", "excused"), ("OD", "leave"))):
        sm.cell(row=r0 + j, column=1, value=rep.t(key)).font = f(bold=True)
        sm.cell(row=r0 + j, column=2, value=f'=COUNTIF(Attendance!$E$2:$E${last},"{code}")').font = f()
    sm.cell(row=r0 + 5, column=1, value=rep.t("total")).font = f(bold=True); sm.cell(row=r0 + 5, column=2, value=f"=SUM(B{r0}:B{r0+4})").font = f()
    sm.cell(row=r0 + 6, column=1, value=rep.t("pct")).font = f(bold=True)
    sm.cell(row=r0 + 6, column=2, value=f'=IFERROR((B{r0}+B{r0+1})/(B{r0+5}-B{r0+3}-B{r0+4}),"")').font = f(); sm.cell(row=r0 + 6, column=2).number_format = "0.0%"
    sm.cell(row=r0 + 8, column=1, value=rep.t("rule")).font = f(italic=True, color="5C6370")
    sm.column_dimensions["A"].width = 22; sm.column_dimensions["B"].width = 30
    bundle = reportstore.load(sid)
    if bundle["images"]:
        ph = wb.create_sheet("Photos")
        row = 1
        for ip in bundle["images"]:
            try:
                im = XLImage(ip); ratio = 900 / im.width; im.width = 900; im.height = int(im.height * ratio)
                ph.add_image(im, f"A{row}"); row += int(im.height / 20) + 3
            except Exception: continue
        ph.cell(row=row, column=1, value=rep.t("legend")).font = f(italic=True)
    ch = wb.create_sheet("Changes"); _xl_head(ch, [rep.t("student"), rep.t("system"), rep.t("final"), rep.t("note")])
    for i, r in enumerate(_change_rows(db, rep, sid, rows) or [[rep.t("none"), "", "", ""]], 2):
        for c, v in enumerate(r, 1): _xl_text(ch.cell(row=i, column=c), v).font = f()
    for i, w in enumerate([36, 16, 16, 50], 1): ch.column_dimensions[get_column_letter(i)].width = w
    wb.save(str(path)); return str(path)


# =========================================================================== DAY
def _day_matrix(db, class_id, date):
    sess = db.sessions_on(class_id, date)
    stu = {}
    for s in sess:
        for r in db.session_rows(s["id"]):
            stu.setdefault(r["code"], {"name": r["name"], "st": {}})["st"][s["id"]] = r["status"]
    return sess, stu


def day_docx(db, class_id, date, path, lang="en"):
    rep = Rep(lang); cname = db.q("SELECT name FROM classes WHERE id=?", (class_id,))[0]["name"]
    sess, stu = _day_matrix(db, class_id, date)
    d = _doc(rep); _head(d, rep, db, rep.t("day_title"))
    _info_table(d, rep, [(rep.t("class"), cname), (rep.t("date"), date)])
    d.add_paragraph()
    heads = [rep.t("id"), rep.t("name")] + [f"{s['time']}\n{s['subject'] or '-'}\n{s['period'] or ''}" for s in sess] + [rep.t("pct")]
    body = []
    for c in sorted(stu):
        sts = [stu[c]["st"].get(s["id"], "") for s in sess]; k = stats.counts([x for x in sts if x])
        body.append([c, stu[c]["name"] or ""] + [x or "-" for x in sts] + ["-" if k["pct"] is None else f"{k['pct']:.0f}%"])
    _table(d, rep, heads, body, colour_fn=lambda row, i: STATUS_HEX.get(row[i]) if 2 <= i < len(row) - 1 and row[i] in STATUS_HEX else None)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    _run(p, "P = " + rep.t("present") + ", A = " + rep.t("absent") + ", L = " + rep.t("late") + ", E = " + rep.t("excused") + ", OD = " + rep.t("leave") + ".   " + rep.t("rule"), rep, italic=True, size=8.5, color="5C6370")
    _signature(d, rep); d.save(str(path)); return str(path)


def day_xlsx(db, class_id, date, path, lang="en"):
    rep = Rep(lang); sess, stu = _day_matrix(db, class_id, date)
    wb = Workbook(); ws = wb.active; ws.title = "Day"; f = lambda **k: Font(name=rep.font, **k)
    _xl_head(ws, [rep.t("id"), rep.t("name")] + [f"{s['time']} {s['subject'] or ''} {s['period'] or ''}".strip() for s in sess] + [rep.t("pct")])
    n = len(sess); lastc = get_column_letter(2 + n)
    for i, c in enumerate(sorted(stu), 2):
        _xl_text(ws.cell(row=i, column=1), c).font = f(); _xl_text(ws.cell(row=i, column=2), stu[c]["name"] or "").font = f()
        for j, s in enumerate(sess):
            v = stu[c]["st"].get(s["id"], ""); x = ws.cell(row=i, column=3 + j, value=v or None); x.font = f(); x.alignment = Alignment(horizontal="center")
            if v: x.fill = PatternFill("solid", start_color=FILL.get(v, "FFFFFF"), end_color=FILL.get(v, "FFFFFF"))
        rng = f"C{i}:{lastc}{i}"
        x = ws.cell(row=i, column=3 + n, value=f'=IFERROR((COUNTIF({rng},"P")+COUNTIF({rng},"L"))/(COUNTA({rng})-COUNTIF({rng},"E")-COUNTIF({rng},"OD")),"")')
        x.number_format = "0%"; x.font = f()
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 26
    for j in range(n + 1): ws.column_dimensions[get_column_letter(3 + j)].width = 16
    ws.freeze_panes = "C2"; ws.row_dimensions[1].height = 34
    wb.save(str(path)); return str(path)


# =========================================================================== MONTH
def month_docx(db, class_id, year, month, path, lang="en"):
    rep = Rep(lang); cname = db.q("SELECT name FROM classes WHERE id=?", (class_id,))[0]["name"]
    rows, nsess = stats.month_students(db, class_id, year, month); warn, lim = stats.limits(db)
    d = _doc(rep); _head(d, rep, db, rep.t("month_title"))
    _info_table(d, rep, [(rep.t("class"), cname), (rep.t("month"), f"{calendar.month_name[month]} {year}"), (rep.t("sessions"), nsess), ("Limits", f"{warn}% / {lim}%")])
    d.add_paragraph()
    lvl = {"ok": rep.t("ok"), "warn": rep.t("warn"), "critical": rep.t("crit"), None: "-"}
    body = [[r["code"], r["name"] or "", r["attended"], r["absent"], r["excused"], r["total"], "-" if r["pct"] is None else f"{r['pct']:.1f}%", lvl[r["level"]]] for r in rows]
    inv = {v: k for k, v in lvl.items()}
    _table(d, rep, [rep.t("id"), rep.t("name"), rep.t("present") + "+" + rep.t("late"), rep.t("absent"), rep.t("excused") + "/OD", rep.t("total"), rep.t("pct"), rep.t("status")], body,
           [1.35, 0.7, 0.95, 0.7, 0.85, 0.55, 0.9, 0.95],
           colour_fn=lambda row, i: LEVEL_HEX.get(inv.get(row[7])) if i in (6, 7) else None)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); _run(p, rep.t("rule"), rep, italic=True, size=8.5, color="5C6370")
    _signature(d, rep); d.save(str(path)); return str(path)


def month_xlsx(db, class_id, year, month, path, lang="en"):
    rep = Rep(lang); sess = db.sessions_in_month(class_id, year, month); warn, lim = stats.limits(db)
    wb = Workbook(); sm = wb.active; sm.title = "Monthly summary"; mx = wb.create_sheet("Sessions")
    f = lambda **k: Font(name=rep.font, **k)
    stu = {}
    for s in sess:
        for r in db.session_rows(s["id"]): stu.setdefault(r["code"], {"name": r["name"], "st": {}})["st"][s["id"]] = r["status"]
    _xl_head(mx, [rep.t("id"), rep.t("name")] + [f"{s['date']} {s['time']} {s['subject'] or ''} {s['period'] or ''}".strip() for s in sess])
    codes = sorted(stu)
    for i, c in enumerate(codes, 2):
        _xl_text(mx.cell(row=i, column=1), c).font = f(); _xl_text(mx.cell(row=i, column=2), stu[c]["name"] or "").font = f()
        for j, s in enumerate(sess):
            v = stu[c]["st"].get(s["id"], ""); x = mx.cell(row=i, column=3 + j, value=v or None); x.alignment = Alignment(horizontal="center"); x.font = f()
            if v: x.fill = PatternFill("solid", start_color=FILL.get(v, "FFFFFF"), end_color=FILL.get(v, "FFFFFF"))
    mx.column_dimensions["A"].width = 18; mx.column_dimensions["B"].width = 26; mx.freeze_panes = "C2"; mx.row_dimensions[1].height = 48
    for j in range(len(sess)): mx.column_dimensions[get_column_letter(3 + j)].width = 15
    sm["A1"] = f"{rep.t('month_title')} - {calendar.month_name[month]} {year}"; sm["A1"].font = f(bold=True, size=14)
    sm["A2"] = "Warning below (%)"; sm["B2"] = warn; sm["A3"] = "Limit (%)"; sm["B3"] = lim
    for c in ("A2", "A3"): sm[c].font = f(bold=True)
    for c in ("B2", "B3"): sm[c].font = Font(name=rep.font, color="0000FF")
    _xl_head(sm, [rep.t("id"), rep.t("name"), rep.t("present"), rep.t("late"), rep.t("absent"), rep.t("excused"), rep.t("leave"), rep.t("total"), rep.t("pct"), rep.t("status")], row=5)
    lastc = get_column_letter(2 + max(1, len(sess)))
    for i, c in enumerate(codes, 6):
        mr = i - 4; rng = f"Sessions!$C{mr}:${lastc}{mr}"
        _xl_text(sm.cell(row=i, column=1), c); _xl_text(sm.cell(row=i, column=2), stu[c]["name"] or "")
        for col, code in ((3, "P"), (4, "L"), (5, "A"), (6, "E"), (7, "OD")):
            sm.cell(row=i, column=col, value=f'=COUNTIF({rng},"{code}")')
        sm.cell(row=i, column=8, value=f"=SUM(C{i}:G{i})")
        sm.cell(row=i, column=9, value=f'=IFERROR((C{i}+D{i})/(H{i}-F{i}-G{i}),"")'); sm.cell(row=i, column=9).number_format = "0.0%"
        sm.cell(row=i, column=10, value=f'=IF(I{i}="","-",IF(I{i}*100<$B$3,"{rep.t("crit")}",IF(I{i}*100<$B$2,"{rep.t("warn")}","{rep.t("ok")}")))')
        for col in range(1, 11): sm.cell(row=i, column=col).font = f()
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    end = 5 + len(codes)
    if codes:
        rngs = f"J6:J{end}"
        sm.conditional_formatting.add(rngs, FormulaRule(formula=[f'$J6="{rep.t("crit")}"'], fill=PatternFill("solid", start_color="FDECEA", end_color="FDECEA"), font=Font(color="C62828", bold=True)))
        sm.conditional_formatting.add(rngs, FormulaRule(formula=[f'$J6="{rep.t("warn")}"'], fill=PatternFill("solid", start_color="FFF3E0", end_color="FFF3E0"), font=Font(color="EF6C00", bold=True)))
    sm.cell(row=end + 2, column=1, value=rep.t("rule")).font = f(italic=True, color="5C6370")
    for i, w in enumerate([18, 26, 9, 9, 9, 10, 10, 8, 12, 16], 1): sm.column_dimensions[get_column_letter(i)].width = w
    sm.freeze_panes = "C6"; wb.save(str(path)); return str(path)


def range_xlsx(db, class_id, start, end, path, lang="en"):
    """Create a class attendance summary workbook for an inclusive ISO date range."""
    # Historical reports include archived students whose attendance rows remain.
    rows, held = stats.range_students(db, class_id, start, end, include_archived=True)
    classes = db.q("SELECT name FROM classes WHERE id=?", (class_id,))
    if not classes: raise ValueError("Class not found.")
    wb = Workbook(); ws = wb.active; ws.title = "Attendance"
    ws.append(["Upasthiti"]); _xl_text(ws.cell(row=1, column=1), db.get("college_name", "") or "Upasthiti")
    ws.append(["Class", classes[0]["name"], "From", start, "To", end]); _xl_text(ws.cell(row=2, column=2), classes[0]["name"])
    ws.append(["Sessions held", held])
    ws.append(["Student ID", "Student name", "Attended", "Absent", "Excused / OD", "Records", "Attendance %"])
    for cell in ws[4]:
        cell.font = Font(bold=True); cell.fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    for r in rows:
        ws.append([r["code"], r["name"] or "", r["attended"], r["absent"], r["excused"], r["total"], None if r["pct"] is None else round(r["pct"], 2)])
        _xl_text(ws.cell(row=ws.max_row, column=1), r["code"]); _xl_text(ws.cell(row=ws.max_row, column=2), r["name"] or "")
    ws.freeze_panes = "A5"
    for col, width in zip("ABCDEFG", (20, 28, 13, 12, 18, 12, 18)):
        ws.column_dimensions[col].width = width
    wb.save(str(path))
    return str(path)
